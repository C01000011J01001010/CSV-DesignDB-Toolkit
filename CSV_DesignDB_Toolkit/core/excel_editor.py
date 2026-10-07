import os
import openpyxl

def get_excel_files(target_dir, recursive):
    # 💡 [FIX] workspace_root 파라미터 제거. 오직 target_dir 기준으로만 상대경로를 뽑습니다.
    files = []
    for root, dirs, fnames in os.walk(target_dir):
        dirs[:] = [d for d in dirs if not d.startswith('Disabled')]
        if not recursive: dirs.clear()
        for f in fnames:
            if f.endswith('.xlsx') and not f.startswith('~$'):
                full_path = os.path.join(root, f)
                try:
                    rel_path = os.path.relpath(full_path, target_dir)
                except ValueError:
                    rel_path = full_path
                files.append(rel_path)
    return files

def read_headers(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        headers = []
        for cell in ws[1]:
            headers.append(str(cell.value) if cell.value is not None else "")
        wb.close()
        return True, headers
    except Exception as e:
        return False, str(e)

def write_headers(filepath, headers):
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h)
        wb.save(filepath)
        wb.close()
        return True, "저장 성공"
    except Exception as e:
        return False, str(e)

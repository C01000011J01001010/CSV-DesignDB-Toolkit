import os
import openpyxl

# 💡 Prefix 적용하여 엑셀 목록 반환
def get_excel_files(target_dir, include_subdirs=True, exclude_prefix="Disabled"):
    excel_files = []
    for root, dirs, files in os.walk(target_dir):
        if any(p.startswith(exclude_prefix) for p in root.replace('\\', '/').split('/')): continue
        if not include_subdirs and root != target_dir: continue
        for file in files:
            if file.startswith('~$') or file.startswith(exclude_prefix): continue
            if file.endswith('.xlsx'):
                rel_path = os.path.relpath(os.path.join(root, file), target_dir)
                excel_files.append(rel_path)
    return excel_files

# 💡 스키마 바운딩 박스 좌표 추적 엔진
def find_schema_bounds(sheet):
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        valid_cols = []
        for col_idx, val in enumerate(row, start=1):
            val_str = str(val).strip() if val is not None else ""
            if val_str.startswith('{') and val_str.endswith('}'):
                valid_cols.append(col_idx)
        if valid_cols:
            return row_idx, valid_cols[0], valid_cols[-1]
    return None, None, None

def read_headers(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.active
        row_idx, start_col, end_col = find_schema_bounds(sheet)
        
        if row_idx is None:
            wb.close()
            return False, "스키마 행을 찾을 수 없습니다."
            
        headers = []
        # {}로 묶인 시작부터 끝 열까지만 가져오기 (외부 메모 무시)
        for col_idx in range(start_col, end_col + 1):
            val = sheet.cell(row=row_idx, column=col_idx).value
            val_str = str(val).strip() if val is not None else ""
            headers.append(val_str)
            
        wb.close()
        return True, headers
    except Exception as e:
        return False, str(e)

def write_headers(filepath, new_headers, orig_offsets=None):
    """
    스키마 헤더 및 컬럼 구조(추가/삭제/이름·타입 변경)를 원본 엑셀에 반영합니다.
    - new_headers: ['{id/int}', '{name/string}', ...] 형태의 헤더 문자열 리스트
    - orig_offsets: 각 새 컬럼이 기존 스키마의 몇 번째 컬럼(0-based)이었는지 나타내는 리스트.
                    새로 추가된 컬럼은 None.
                    지정하지 않으면(예: 1.2 제약조건 탭 저장) 기존 컬럼 순서(0, 1, 2...) 그대로 유지.
    """
    try:
        if not new_headers:
            return False, "최소 1개 이상의 유효한 스키마 컬럼이 필요합니다."

        wb = openpyxl.load_workbook(filepath)
        sheet = wb.active
        row_idx, start_col, end_col = find_schema_bounds(sheet)
        
        if row_idx is None:
            wb.close()
            return False, "원본 스키마 행을 찾을 수 없습니다."

        old_num_cols = (end_col - start_col + 1)
        new_num_cols = len(new_headers)
        max_r = sheet.max_row

        if orig_offsets is None:
            orig_offsets = [i if i < old_num_cols else None for i in range(new_num_cols)]

        # 컬럼 삭제 또는 순서/구조 변경이 발생했는지 확인
        is_same_positions = (
            new_num_cols >= old_num_cols
            and all(orig_offsets[i] == i for i in range(old_num_cols))
            and all(orig_offsets[i] is None for i in range(old_num_cols, new_num_cols))
        )

        if is_same_positions:
            # 기존 컬럼 위치가 그대로이고 끝에 새 컬럼만 추가(또는 동일 개수 수정)된 경우:
            # 기존 데이터 셀은 건드리지 않고 헤더만 start_col + i 에 제한 없이 기록
            for i, header in enumerate(new_headers):
                c_idx = start_col + i
                sheet.cell(row=row_idx, column=c_idx).value = header
        else:
            # 컬럼 삭제가 포함되었거나 위치가 바뀐 경우:
            # 데이터 행(row_idx + 1 ~ max_r)의 값을 기존 컬럼 오프셋 기준으로 백업 후 재배치
            old_col_data = {}
            for offset in range(old_num_cols):
                c_idx = start_col + offset
                old_col_data[offset] = [
                    sheet.cell(row=r, column=c_idx).value
                    for r in range(row_idx + 1, max_r + 1)
                ]

            for new_i, header in enumerate(new_headers):
                target_c_idx = start_col + new_i
                sheet.cell(row=row_idx, column=target_c_idx).value = header

                src_offset = orig_offsets[new_i] if new_i < len(orig_offsets) else None
                if src_offset is not None and src_offset in old_col_data:
                    col_vals = old_col_data[src_offset]
                    for r_offset, val in enumerate(col_vals):
                        sheet.cell(row=row_idx + 1 + r_offset, column=target_c_idx).value = val
                else:
                    for r in range(row_idx + 1, max_r + 1):
                        sheet.cell(row=r, column=target_c_idx).value = None

            # 삭제되어 줄어든 우측 잔여 컬럼 영역(헤더 + 데이터) 비우기
            if new_num_cols < old_num_cols:
                for leftover_i in range(new_num_cols, old_num_cols):
                    leftover_c_idx = start_col + leftover_i
                    for r in range(row_idx, max_r + 1):
                        sheet.cell(row=r, column=leftover_c_idx).value = None

        wb.save(filepath)
        wb.close()
        return True, ""
    except PermissionError:
        return False, "엑셀 파일이 현재 열려 있어 저장할 수 없습니다. 엑셀 창을 닫고 다시 시도해주세요."
    except Exception as e:
        return False, str(e)

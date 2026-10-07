import os
import openpyxl

def get_excel_files(target_dir, recursive):
    files = []
    for root, dirs, fnames in os.walk(target_dir):
        dirs[:] = [d for d in dirs if not d.startswith('Disabled')]
        if not recursive: dirs.clear()
        for f in fnames:
            if f.endswith('.xlsx') and not f.startswith('~$'):
                full_path = os.path.join(root, f)
                # 루트 경로(target_dir)를 제외한 상대 경로만 추출
                rel_path = os.path.relpath(full_path, target_dir)
                files.append(rel_path)
    return files

def read_headers(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        headers = []
        for cell in ws[1]:
            headers.append(str(cell.value) if cell.value is not None else "")
        wb.close()
        return True, headers
    except Exception as e:
        return False, str(e)

def write_headers(filepath, headers):
    try:
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        for i, h in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=h)
        wb.save(filepath)
        wb.close()
        return True, "저장 성공"
    except Exception as e:
        return False, str(e)
